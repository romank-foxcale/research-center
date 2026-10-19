"""Turn one original file into committed text: text files as they are, documents converted to Markdown.

Originals are never committed (docs/decisions/0002-committed-extracted-sources.md). Every converter is a Python
library locked in uv.lock - pypdfium2, mammoth, markdownify, striprtf, python-pptx, openpyxl - so extraction needs
no system binary and works the same on Windows and POSIX (docs/decisions/0012-windows-native-uv.md).
"""

from __future__ import annotations

import io
import posixpath
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

TEXT_SUFFIXES = {
    ".md",
    ".markdown",
    ".txt",
    ".rst",
    ".adoc",
    ".org",
    ".tex",
    ".bib",
    ".csv",
    ".tsv",
    ".json",
    ".jsonl",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".cfg",
    ".conf",
    ".xml",
    ".sql",
    ".graphql",
    ".proto",
    ".nix",
    ".sh",
    ".bash",
    ".zsh",
    ".ps1",
    ".bat",
    ".cmd",
    ".py",
    ".ipynb",
    ".r",
    ".jl",
    ".js",
    ".mjs",
    ".cjs",
    ".ts",
    ".tsx",
    ".jsx",
    ".vue",
    ".svelte",
    ".css",
    ".scss",
    ".java",
    ".kt",
    ".kts",
    ".scala",
    ".go",
    ".rs",
    ".c",
    ".h",
    ".cc",
    ".cpp",
    ".hpp",
    ".cs",
    ".fs",
    ".swift",
    ".m",
    ".mm",
    ".rb",
    ".php",
    ".pl",
    ".lua",
    ".dart",
    ".ex",
    ".exs",
    ".erl",
    ".hs",
    ".clj",
    ".elm",
    ".zig",
    ".tf",
    ".hcl",
    ".cmake",
    ".gradle",
    ".properties",
    ".env.example",
    ".dockerfile",
}
TEXT_NAMES = {"dockerfile", "makefile", "justfile", "license", "readme", "changelog", "contributing", "codeowners"}
SKIPPED_NAMES = {
    "package-lock.json",
    "yarn.lock",
    "pnpm-lock.yaml",
    "cargo.lock",
    "flake.lock",
    "poetry.lock",
    "uv.lock",
    "composer.lock",
    "gemfile.lock",
    "go.sum",
}
UNSEEN = ("script", "style", "noscript", "template", "svg", "iframe")
CHARSET = re.compile(rb"""<meta[^>]+charset=["']?([\w-]+)""", re.IGNORECASE)


class Skip(Exception):
    """The file is not ingested; the message is the reason recorded in the manifest."""


def stored_name(original: str, converted: bool) -> str:
    """Where the text of an original lands: Markdown stays .md, everything else gets a suffix, so no
    ingested file can act on Git or the tools (.gitignore, .gitattributes, pyproject.toml). Characters outside
    [A-Za-z0-9._/-] become '_', so every stored path can be cited as [sources/...] without quoting."""
    safe = re.sub(r"[^A-Za-z0-9._/-]", "_", original)
    if not converted and safe.lower().endswith(".md"):
        return safe
    return f"{safe}.md" if converted else f"{safe}.txt"


def is_text_name(name: str) -> bool:
    lower = Path(name).name.lower()
    suffix = Path(lower).suffix
    return suffix in TEXT_SUFFIXES or lower in TEXT_NAMES or Path(lower).stem in TEXT_NAMES


def decode(data: bytes) -> str:
    if b"\0" in data[:8192]:
        raise Skip("binary content")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def safe_xml(data: bytes) -> ElementTree.Element:
    """Parse XML from a document. Entity declarations are refused: they are how XML expands into a memory bomb."""
    if b"<!ENTITY" in data:
        raise Skip("the document's XML declares entities; refused")
    try:
        return ElementTree.fromstring(data)
    except ElementTree.ParseError as error:
        raise Skip(f"malformed XML in the document: {error}") from error


# Converters ---------------------------------------------------------------------------------------------------


def html_to_markdown(data: bytes | str) -> str:
    """Markdown from HTML; scripts, styles and other unseen parts dropped."""
    from bs4 import BeautifulSoup  # markdownify's own dependency
    from markdownify import markdownify

    if isinstance(data, bytes):
        declared = CHARSET.search(data[:4096])
        try:
            data = data.decode(declared.group(1).decode("ascii") if declared else "utf-8")
        except (LookupError, UnicodeDecodeError):
            data = data.decode("utf-8", errors="replace")
    soup = BeautifulSoup(data, "html.parser")
    for element in soup(UNSEEN):
        element.decompose()
    text = markdownify(str(soup), heading_style="ATX", bullets="-")
    return re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"


def pdf_text(data: bytes) -> str:
    import pypdfium2

    try:
        document = pypdfium2.PdfDocument(data)
    except pypdfium2.PdfiumError as error:
        raise Skip(f"unreadable PDF: {error}") from error
    pages = []
    try:
        for number in range(len(document)):
            page = document[number]
            text = page.get_textpage().get_text_range().replace("\r\n", "\n").replace("\r", "\n")
            pages.append(f"## Page {number + 1}\n\n{text.strip()}")
    finally:
        document.close()
    if not any(part.split("\n", 2)[-1].strip() for part in pages):
        raise Skip("the PDF has no text layer (scanned images need OCR)")
    return "\n\n".join(pages) + "\n"


def docx_text(data: bytes) -> str:
    import mammoth

    return html_to_markdown(mammoth.convert_to_html(io.BytesIO(data)).value)


def epub_text(data: bytes) -> str:
    """The chapters of an EPUB in reading order (the OPF spine), each converted from XHTML."""
    with zipfile.ZipFile(io.BytesIO(data)) as book:
        container = safe_xml(book.read("META-INF/container.xml"))
        rootfile = next(
            (element.get("full-path") for element in container.iter() if element.tag.endswith("rootfile")), None
        )
        if not rootfile:
            raise Skip("EPUB without a package file")
        package = safe_xml(book.read(rootfile))
        base = posixpath.dirname(rootfile)
        manifest = {item.get("id"): item.get("href") for item in package.iter() if item.tag.endswith("}item")}
        spine = [item.get("idref") for item in package.iter() if item.tag.endswith("}itemref")]
        chapters = []
        for reference in spine:
            href = manifest.get(reference)
            if href:
                chapters.append(html_to_markdown(book.read(posixpath.normpath(posixpath.join(base, href)))).strip())
    return "\n\n".join(chapter for chapter in chapters if chapter) + "\n"


def odt_text(data: bytes) -> str:
    """Headings and paragraphs of an OpenDocument text, in order."""
    with zipfile.ZipFile(io.BytesIO(data)) as document:
        root = safe_xml(document.read("content.xml"))
    lines = []
    for element in root.iter():
        tag = element.tag.rsplit("}", 1)[-1]
        if tag in ("h", "p"):
            text = " ".join("".join(element.itertext()).split())
            if text:
                level = next((int(value) for key, value in element.attrib.items() if key.endswith("outline-level")), 1)
                lines.append(f"{'#' * min(level, 6)} {text}" if tag == "h" else text)
    return "\n\n".join(lines) + "\n"


def rtf_text(data: bytes) -> str:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(data.decode("latin-1"), errors="ignore").strip() + "\n"


def pptx_text(data: bytes) -> str:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    parts = []
    for number, slide in enumerate(deck.slides, 1):
        texts = [
            shape.text_frame.text.strip()
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        parts.append(f"## Slide {number}\n\n" + "\n\n".join(texts) + (f"\n\nNotes: {notes}" if notes else ""))
    return "\n\n".join(parts)


def xlsx_text(data: bytes, max_rows: int = 200) -> str:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in book.worksheets:
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= max_rows:
                rows.append(f"(first {max_rows} rows only)")
                break
            rows.append(" | ".join("" if cell is None else str(cell) for cell in row))
        parts.append(f"## Sheet {sheet.title}\n\n" + "\n".join(rows))
    return "\n\n".join(parts)


CONVERTERS = {
    ".pdf": (pdf_text, "pypdfium2"),
    ".docx": (docx_text, "mammoth"),
    ".epub": (epub_text, "epub"),
    ".odt": (odt_text, "odt"),
    ".rtf": (rtf_text, "striprtf"),
    ".html": (html_to_markdown, "markdownify"),
    ".htm": (html_to_markdown, "markdownify"),
    ".pptx": (pptx_text, "python-pptx"),
    ".xlsx": (xlsx_text, "openpyxl"),
}


def extract(name: str, data: bytes) -> tuple[str, str, str]:
    """(text, stored name, extractor) for one original file, or Skip."""
    lower = Path(name).name.lower()
    suffix = Path(lower).suffix
    if lower in SKIPPED_NAMES or lower.endswith((".min.js", ".min.css", ".map")):
        raise Skip("generated or lock file")
    if suffix in CONVERTERS:
        convert, extractor = CONVERTERS[suffix]
        try:
            return convert(data), stored_name(name, True), extractor
        except Skip:
            raise
        except (zipfile.BadZipFile, KeyError, ValueError, OSError) as error:
            raise Skip(f"{extractor} could not read it: {str(error)[:160]}") from error
    if is_text_name(name):
        return decode(data), stored_name(name, False), "text"
    raise Skip(f"unsupported file type {suffix or lower}")
