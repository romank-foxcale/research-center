#!/usr/bin/env python3
"""Test extraction of every document type with no system binary: the test builds a small PDF, DOCX, EPUB, ODT,
RTF, HTML, PPTX and XLSX itself and checks the text, the stored name and the extractor of each, plus the refusals
(scanned PDF, entity bomb, corrupt archive)."""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import extract  # noqa: E402


def pdf(pages: list[str]) -> bytes:
    """A minimal PDF with one line of Helvetica text per page, offsets computed so the xref table is valid."""
    count = len(pages)
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(f'{3 + 2 * index} 0 R' for index in range(count))}] /Count {count} >>",
    ]
    font = 3 + 2 * count
    for index, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * index} 0 R "
            f"/Resources << /Font << /F1 {font} 0 R >> >> >>"
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1"))
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        out.write(f"{offset:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def archive(files: dict[str, str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as handle:
        for name, content in files.items():
            handle.writestr(name, content)
    return buffer.getvalue()


def docx() -> bytes:
    word = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    return archive(
        {
            "[Content_Types].xml": '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" '
            'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>',
            "_rels/.rels": '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
            'Target="word/document.xml"/></Relationships>',
            "word/document.xml": f'<?xml version="1.0"?><w:document xmlns:w="{word}"><w:body>'
            '<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Retention policy</w:t></w:r></w:p>'
            "<w:p><w:r><w:t>Mail is kept for </w:t></w:r><w:r><w:rPr><w:b/></w:rPr><w:t>seven years</w:t></w:r></w:p>"
            "<w:p><w:r><w:t>保存期間は七年です</w:t></w:r></w:p></w:body></w:document>",
        }
    )


def epub() -> bytes:
    return archive(
        {
            "mimetype": "application/epub+zip",
            "META-INF/container.xml": '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="OEBPS/book.opf" media-type="application/oebps-package+xml"/></rootfiles></container>',
            "OEBPS/book.opf": '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0"><manifest>'
            '<item id="c2" href="text/two.xhtml" media-type="application/xhtml+xml"/>'
            '<item id="c1" href="text/one.xhtml" media-type="application/xhtml+xml"/></manifest>'
            '<spine><itemref idref="c1"/><itemref idref="c2"/></spine></package>',
            "OEBPS/text/one.xhtml": "<html><body><h1>Chapter One</h1><p>The first finding.</p></body></html>",
            "OEBPS/text/two.xhtml": "<html><body><h1>Chapter Two</h1><p>The second finding.</p></body></html>",
        }
    )


def odt() -> bytes:
    return archive(
        {
            "content.xml": '<?xml version="1.0"?><office:document-content '
            'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
            '<text:h text:outline-level="2">Findings</text:h><text:p>Output held <text:span>steady</text:span>.</text:p>'
            "</office:text></office:body></office:document-content>"
        }
    )


def expect(name: str, data: bytes, contains: list[str], extractor: str, absent: tuple[str, ...] = ()) -> str:
    text, stored, used = extract.extract(name, data)
    assert stored == f"{name}.md" and used == extractor, (name, stored, used)
    missing = [part for part in contains if part not in text]
    assert not missing and not [part for part in absent if part in text], (name, missing, text)
    return text


def refused(name: str, data: bytes, reason: str) -> None:
    try:
        extract.extract(name, data)
    except extract.Skip as error:
        assert reason in str(error), (name, str(error))
        return
    raise AssertionError(f"{name} must be skipped")


def main() -> int:
    text = expect(
        "report.pdf",
        pdf(["Revenue rose 1.4 percent", "Output held"]),
        ["## Page 1", "Revenue rose 1.4 percent", "## Page 2", "Output held"],
        "pypdfium2",
    )
    assert text.index("Page 1") < text.index("Page 2")
    expect("policy.docx", docx(), ["# Retention policy", "seven years", "保存期間は七年です"], "mammoth")
    text = expect("book.epub", epub(), ["# Chapter One", "The first finding.", "# Chapter Two"], "epub")
    assert text.index("Chapter One") < text.index("Chapter Two"), "the spine decides the order, not the manifest"
    expect("notes.odt", odt(), ["## Findings", "Output held steady."], "odt")
    expect(
        "memo.rtf",
        rb"{\rtf1\ansi{\fonttbl\f0\fswiss Helvetica;}\f0\pard Trial ended \b early\b0.\par}",
        ["Trial ended early."],
        "striprtf",
    )
    expect(
        "page.html",
        "<html><head><meta charset='shift_jis'><style>p{}</style></head><body><h2>結果</h2>"
        "<p>生産性は維持</p><script>track()</script></body></html>".encode("shift_jis"),
        ["## 結果", "生産性は維持"],
        "markdownify",
        absent=("track()", "p{}"),
    )

    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Pilot design"
    slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "Six months, one team"
    buffer = io.BytesIO()
    deck.save(buffer)
    expect("deck.pptx", buffer.getvalue(), ["## Slide 1", "Pilot design", "Six months, one team"], "python-pptx")
    book = Workbook()
    book.active.append(["metric", "baseline"])
    book.active.append(["overtime", 12])
    buffer = io.BytesIO()
    book.save(buffer)
    expect("metrics.xlsx", buffer.getvalue(), ["metric | baseline", "overtime | 12"], "openpyxl")

    refused("scan.pdf", pdf([""]), "no text layer")
    refused(
        "bomb.odt",
        archive({"content.xml": '<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><x>&a;</x>'}),
        "entities",
    )
    refused("broken.docx", b"PK\x03\x04 not really a zip", "could not read it")
    refused("broken.pdf", b"%PDF-1.4 garbage", "PDF")
    print("extract test passed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
